import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

class ExcelProcessor:
    def __init__(self, sheets_dict, name_checker, company_code, cwid_input, street1_remove, state_cleanup,
                timezone_input, phone_number, aci_airport, relationship, station_so, postal_canada,some_conditions):
        self.sheets_dict = sheets_dict
        self.name_checker = name_checker
        self.company_code = company_code
        self.cwid_input = cwid_input
        self.street1_remove = street1_remove
        self.state_cleanup = state_cleanup
        self.timezone_input = timezone_input
        # self.company_name = company_name
        self.phone_number = phone_number
        self.aci_airport = aci_airport
        self.relationship = relationship
        self.station_so = station_so
        self.postal_canada = postal_canada
        self.some_conditions = some_conditions

    def process_names(self, sheet_name, start_row, end_row, name_column, output_column):
        """Process the names in the given Excel file."""
        if sheet_name not in self.sheets_dict:
            raise ValueError(f"Sheet {sheet_name} does not exist in the provided sheets.")

        df = self.sheets_dict[sheet_name]  # Get the DataFrame for the specific sheet

        # Ensure column names are lowercase for consistent access
        df.columns = df.columns.str.lower()
        name_column = name_column.lower()
        output_column = output_column.lower()

        if name_column not in df.columns or output_column not in df.columns:
            raise ValueError(f"Columns '{name_column}' or '{output_column}' do not exist in the sheet.")

        # Iterate over the rows and process names
        for index, row in df.iterrows():
            if start_row <= index <= end_row:
                name = row.get(name_column, None)  # Get the name from the name_column
                result = self.name_checker.check_name(name)
                df.at[index, output_column] = result  # Update the output column with the result

        # Update the sheet in sheets_dict with the modified DataFrame
        self.sheets_dict[sheet_name] = df
        return self.sheets_dict

    def save_and_autofit(self, output_file):
        # Save each DataFrame in sheets_dict to an Excel sheet
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            for sheet_name, df in self.sheets_dict.items():
                df.to_excel(writer, sheet_name=sheet_name, index=False)

        # Open the workbook and autofit each column
        workbook = load_workbook(output_file)
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            for column_cells in worksheet.columns:
                max_length = 0
                column_letter = get_column_letter(column_cells[0].column)
                for cell in column_cells:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except:
                        pass
                worksheet.column_dimensions[column_letter].width = max_length + 2  # Adjust width with padding

        # Save the workbook with autofitted columns
        workbook.save(output_file)
